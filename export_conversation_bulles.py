# -*- coding: utf-8 -*-
"""
Export conversation Excel -> Word (bulles) — LibreOffice V5 (CustomTkinter)
Repatch UX + medias + page options

Dépendances optionnelles:
- Pillow (PIL) pour compresser les images : pip install pillow
- ffmpeg dans le PATH pour miniatures vidéo

Notes:
- Les réglages "médias" sont dans l'onglet Fichier (bêta)
- L'affichage date/heure est activé par défaut (pas de case)
- Statuts (effacé/corbeille...) déplacés en Base
"""

import os
import re
import sys
import subprocess
import tempfile
import unicodedata
import json
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, colorchooser, ttk
from typing import Optional

import pandas as pd
import customtkinter as ctk

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENTATION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SETTINGS_PATH = Path.home() / ".export_conversation_bulles_libreoffice_settings.json"


# -----------------------------
# Couleurs utilitaires
# -----------------------------
RECIPIENT_PALETTE = [
    "#7ED957", "#3D7DFF", "#FFD166", "#EF476F", "#06D6A0", "#118AB2",
    "#9B5DE5", "#F15BB5", "#00BBF9", "#00F5D4", "#FEE440", "#FB5607",
]


def split_recipients(v: str) -> list[str]:
    """Découpe une liste de destinataires (To/Participants) en tokens."""
    if v is None:
        return []
    s = str(v).strip()
    if not s:
        return []
    parts = re.split(r"[;\|\n\r,]+", s)
    out: list[str] = []
    for p in parts:
        p = p.strip()
        if p:
            out.append(p)
    return out


def _pick_recipient_color_map(unique_recipients: list[str]) -> dict:
    """Attribue une couleur (hex) à chaque destinataire, de manière déterministe."""
    cmap = {}
    for i, rcp in enumerate(unique_recipients):
        cmap[rcp] = RECIPIENT_PALETTE[i % len(RECIPIENT_PALETTE)]
    return cmap


# -----------------------------
# Excel loading helpers
# -----------------------------
COMMON_HEADER_TOKENS = {
    "from", "to", "body", "direction", "timestamp", "date", "time", "#",
    "attachment", "media", "path", "sender", "deleted", "trash", "subject"
}


def looks_like_hidden_header(first_row_values) -> bool:
    tokens = {str(v).strip().lower() for v in first_row_values if pd.notna(v)}
    hits = sum(1 for t in tokens if any(k in t for k in COMMON_HEADER_TOKENS))
    return hits >= 3


def list_sheets(path: str):
    try:
        xl = pd.ExcelFile(path)
        return xl.sheet_names
    except Exception:
        return []


def load_excel_smart(path: str, sheet_name=None) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=sheet_name)
    if isinstance(df, dict):
        df = df[list(df.keys())[0]]

    if df.shape[0] > 1 and looks_like_hidden_header(df.iloc[0].tolist()):
        header_row = df.iloc[0].tolist()
        real_cols = []
        for i, v in enumerate(header_row):
            if pd.isna(v) or str(v).strip() == "":
                real_cols.append(f"col_{i}")
            else:
                real_cols.append(str(v).strip())
        df = df.copy()
        df.columns = real_cols
        df = df.drop(index=df.index[0]).reset_index(drop=True)

    return df


# -----------------------------
# Word formatting helpers
# -----------------------------
def _hex6(color_hex: str) -> str:
    if not color_hex:
        return "FFFFFF"
    s = color_hex.strip()
    if s.startswith("#"):
        s = s[1:]
    if len(s) == 3:
        s = "".join([c * 2 for c in s])
    return s.upper()


def set_cell_shading(cell, fill_hex: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex6(fill_hex))
    tc_pr.append(shd)


def set_cell_border_none(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "nil")
        tcBorders.append(elem)
    tcPr.append(tcBorders)


def set_cell_width_cm(cell, width_cm: float):
    """Force une largeur de cellule (en cm) via w:tcW pour éviter que Word ignore cell.width."""
    try:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcW = OxmlElement("w:tcW")
        tcW.set(qn("w:type"), "dxa")
        tcW.set(qn("w:w"), str(int(Cm(width_cm).twips)))
        tcPr.append(tcW)
    except Exception:
        pass


def set_table_fixed_layout(table):
    """Fixe le tableau en layout 'fixed' (Word), indispensable pour respecter les largeurs."""
    try:
        table.autofit = False
        tbl = table._tbl
        tblPr = tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement("w:tblPr")
            tbl.append(tblPr)
        tblLayout = OxmlElement("w:tblLayout")
        tblLayout.set(qn("w:type"), "fixed")
        tblPr.append(tblLayout)
    except Exception:
        pass





def set_table_width_cm(table, width_cm: float):
    """Largeur totale explicite du tableau. LibreOffice respecte mieux les largeurs avec tblW + tcW."""
    try:
        tbl = table._tbl
        tblPr = tbl.tblPr
        if tblPr is None:
            tblPr = OxmlElement("w:tblPr")
            tbl.append(tblPr)
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:type"), "dxa")
        tblW.set(qn("w:w"), str(int(Cm(width_cm).twips)))
    except Exception:
        pass


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    """Marges internes de cellule en twips. Compatible Word et LibreOffice."""
    try:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcMar = tcPr.find(qn("w:tcMar"))
        if tcMar is None:
            tcMar = OxmlElement("w:tcMar")
            tcPr.append(tcMar)
        for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
            node = tcMar.find(qn(f"w:{m}"))
            if node is None:
                node = OxmlElement(f"w:{m}")
                tcMar.append(node)
            node.set(qn("w:w"), str(int(v)))
            node.set(qn("w:type"), "dxa")
    except Exception:
        pass

def set_table_borders_none(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.append(tblPr)

    tblBorders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "nil")
        tblBorders.append(elem)
    tblPr.append(tblBorders)


def set_row_cant_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def keep_paragraph_together(paragraph):
    fmt = paragraph.paragraph_format
    fmt.keep_together = True
    fmt.keep_with_next = False


def compact_paragraph(paragraph):
    """Supprime les espacements parasites, surtout visibles dans LibreOffice."""
    try:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
    except Exception:
        pass


def clean_multiline_text(text: str) -> str:
    """Conserve les retours utiles mais supprime les lignes vides qui créent des sauts dans les bulles."""
    if text is None:
        return ""
    text = str(text).replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.strip() for line in text.split("\n")]
    lines = [line for line in lines if line != ""]
    return "\n".join(lines).strip()


def add_text_with_linebreaks(
    paragraph,
    text: str,
    font_size_pt=10,
    bold=False,
    italic=False,
    font_name="Calibri",
    font_color_hex: str | None = None,
):
    compact_paragraph(paragraph)
    text = clean_multiline_text(text)
    lines = text.splitlines() or [""]
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(font_size_pt)
        run.font.name = font_name
        if font_color_hex:
            try:
                run.font.color.rgb = RGBColor.from_string(_hex6(font_color_hex))
            except Exception:
                pass
        if i < len(lines) - 1:
            run.add_break()


def normalize_direction(val: str) -> str:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ""
    s = str(val).strip().lower()
    if "out" in s:
        return "outgoing"
    if "in" in s:
        return "incoming"
    return s


def excel_serial_to_datetime_string(x) -> str:
    try:
        if x is None or (isinstance(x, float) and pd.isna(x)):
            return ""
        if isinstance(x, (int, float)) and not isinstance(x, bool):
            dt = pd.to_datetime(x, unit="D", origin="1899-12-30", errors="coerce")
            if pd.isna(dt):
                return str(x)
            if dt.hour == 0 and dt.minute == 0 and dt.second == 0:
                return dt.strftime("%d/%m/%Y")
            return dt.strftime("%d/%m/%Y %H:%M:%S")
        if isinstance(x, pd.Timestamp):
            if pd.isna(x):
                return ""
            return x.strftime("%d/%m/%Y %H:%M:%S")
        return str(x).strip()
    except Exception:
        return str(x).strip()


def excel_time_to_string(x) -> str:
    """Convertit proprement une heure Excel isolée, sans date parasite 30/12/1899."""
    try:
        if x is None or (isinstance(x, float) and pd.isna(x)):
            return ""
        if isinstance(x, (int, float)) and not isinstance(x, bool):
            if 0 <= float(x) < 1:
                total_seconds = int(round(float(x) * 24 * 60 * 60)) % (24 * 60 * 60)
                h = total_seconds // 3600
                m = (total_seconds % 3600) // 60
                sec = total_seconds % 60
                return f"{h:02d}:{m:02d}" if sec == 0 else f"{h:02d}:{m:02d}:{sec:02d}"
            return excel_serial_to_datetime_string(x)
        if isinstance(x, pd.Timestamp):
            if pd.isna(x):
                return ""
            return x.strftime("%H:%M") if x.second == 0 else x.strftime("%H:%M:%S")
        # datetime.time sans importer explicitement datetime
        if hasattr(x, "hour") and hasattr(x, "minute") and not hasattr(x, "date"):
            sec = getattr(x, "second", 0)
            return f"{x.hour:02d}:{x.minute:02d}" if sec == 0 else f"{x.hour:02d}:{x.minute:02d}:{sec:02d}"
        return str(x).strip()
    except Exception:
        return str(x).strip()


def build_timestamp(row, ts_single_col, ts_date_col, ts_time_col) -> str:
    def safe(v):
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return ""
        return v

    if ts_single_col and ts_single_col in row:
        v = safe(row[ts_single_col])
        if isinstance(v, str):
            return v.strip()
        return excel_serial_to_datetime_string(v)

    d = safe(row.get(ts_date_col, "")) if ts_date_col else ""
    t = safe(row.get(ts_time_col, "")) if ts_time_col else ""

    if isinstance(t, str) and re.search(r"\d{2}/\d{2}/\d{4}", t):
        return t.strip()

    d_str = excel_serial_to_datetime_string(d) if d != "" else ""
    t_str = excel_time_to_string(t) if t != "" else ""

    if d_str and t_str:
        if re.search(r"\d{2}/\d{2}/\d{4}", t_str):
            return t_str
        return f"{d_str} {t_str}"
    return d_str or t_str


def is_image_file(path: str) -> bool:
    ext = os.path.splitext(path)[1].lower()
    return ext in {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff"}


def is_video_file(path: str) -> bool:
    ext = os.path.splitext(path)[1].lower()
    return ext in {".mp4", ".mov", ".avi", ".mkv", ".wmv", ".webm", ".m4v"}


def resolve_status_label(row, status_col: str | None, status_map: dict) -> str:
    if not status_col or status_col not in row or not status_map:
        return ""
    v = row.get(status_col, "")
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return status_map.get(str(v).strip(), "")


def add_page_number(section, fmt="PAGE / NUMPAGES"):
    """Add page numbering in footer: 'PAGE / NUMPAGES' -> e.g. '1 / 2'."""
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _add_field(run, instr_text: str):
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")

        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = f" {instr_text} "

        fld_separate = OxmlElement("w:fldChar")
        fld_separate.set(qn("w:fldCharType"), "separate")

        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")

        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_separate)
        run._r.append(fld_end)

    run = p.add_run()
    _add_field(run, "PAGE")
    run.add_text(" / ")
    _add_field(run, "NUMPAGES")


# -----------------------------
# Dedup helpers (Option B)
# -----------------------------
def _normalize_body(text: str) -> str:
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return ""
    s = str(text)
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    s = "\n".join(line.rstrip() for line in s.split("\n"))
    s = s.strip()
    s = re.sub(r"[ \t]+", " ", s)
    return s


def _normalize_participant(val) -> str:
    """Normalise une valeur participant (robuste aux NBSP/espaces invisibles)."""
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return ""
    s = str(val)
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("\u00A0", " ")  # NBSP
    s = s.replace("\u202F", " ")  # narrow no-break space
    s = s.replace("\u2007", " ")  # figure space
    s = s.replace("\u200B", "")   # zero-width space
    s = s.replace("\uFEFF", "")   # BOM / zero-width no-break
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def dedupe_messages_option_b(df: pd.DataFrame, col_sender: str, col_message: str, ts_series: pd.Series) -> pd.DataFrame:
    tmp = df.copy()
    tmp["_from_norm"] = tmp[col_sender].fillna("").astype(str).str.strip()
    tmp["_body_norm"] = tmp[col_message].apply(_normalize_body) if col_message in tmp.columns else ""
    tmp["_ts_norm"] = ts_series.fillna("").astype(str).str.strip()
    tmp = tmp.drop_duplicates(subset=["_from_norm", "_body_norm", "_ts_norm"], keep="first")
    return tmp.drop(columns=["_from_norm", "_body_norm", "_ts_norm"], errors="ignore")


# -----------------------------
# Open Word after export
# -----------------------------
def open_in_word(path: str):
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.run(["open", path], check=False)
        else:
            subprocess.run(["xdg-open", path], check=False)
    except Exception:
        pass


# -----------------------------
# Media helpers
# -----------------------------
def split_media_values(v: str) -> list[str]:
    if not v:
        return []
    s = str(v).strip()
    if not s:
        return []
    parts = re.split(r"[;\|\n\r,]+", s)
    return [p.strip() for p in parts if p.strip()]


def resolve_media_path(name: str, root_dir: str, recursive: bool = True, fuzzy: bool = True) -> str | None:
    if not name:
        return None
    name = str(name).strip()
    if not name:
        return None

    p = Path(name)
    if p.is_absolute() and p.exists():
        return str(p)

    if root_dir and os.path.isdir(root_dir):
        direct = Path(root_dir) / name
        if direct.exists():
            return str(direct)

        if recursive:
            for root, _, files in os.walk(root_dir):
                if name in files:
                    return str(Path(root) / name)

            if fuzzy:
                low = name.lower()
                for root, _, files in os.walk(root_dir):
                    for f in files:
                        if low in f.lower():
                            return str(Path(root) / f)

    return None


def pillow_available() -> bool:
    try:
        import PIL  # noqa: F401
        return True
    except Exception:
        return False


def compress_image_for_docx(src_path: str, max_px: int, quality: int, cache_dir: str) -> str:
    if not pillow_available():
        return src_path
    try:
        from PIL import Image
    except Exception:
        return src_path

    src = Path(src_path)
    if not src.exists():
        return src_path

    out_dir = Path(cache_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / (src.stem + f"_q{int(quality)}_px{int(max_px)}.jpg")

    try:
        with Image.open(str(src)) as im:
            im = im.convert("RGB")
            w, h = im.size
            longest = max(w, h)
            if longest > max_px:
                scale = max_px / float(longest)
                new_size = (max(1, int(w * scale)), max(1, int(h * scale)))
                im = im.resize(new_size)
            im.save(str(out_path), format="JPEG", quality=int(quality), optimize=True)
        return str(out_path)
    except Exception:
        return src_path


def ffmpeg_available() -> bool:
    try:
        subprocess.run(["ffmpeg", "-version"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        return True
    except Exception:
        return False


def make_video_thumbnail(video_path: str, cache_dir: str) -> str | None:
    src = Path(video_path)
    if not src.exists():
        return None
    if not ffmpeg_available():
        return None

    out_dir = Path(cache_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / (src.stem + "_thumb.png")
    cmd = ["ffmpeg", "-y", "-ss", "00:00:01", "-i", str(src), "-frames:v", "1", str(out_path)]
    try:
        subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        return str(out_path) if out_path.exists() else None
    except Exception:
        return None


# -----------------------------
# DOCX exporter
# -----------------------------
def export_chat_to_docx(
    df: pd.DataFrame,
    out_path: str,
    col_sender: str,
    col_message: str,
    col_direction: str | None,
    mode_side: str,  # "direction" or "sender"
    right_participant_value: str | None,
    owner_col: str | None,
    ts_single_col: str | None,
    ts_date_col: str | None,
    ts_time_col: str | None,
    display_name_map: dict,
    empty_sender_display: str,
    participant_color_map: dict,
    default_left_color: str,
    default_right_color: str,
    bubble_width_cm: float,
    gutter_cm: float,
    bubble_spacing_pt: float,
    extra_cols: Optional[list],
    media_path_col: str | None,
    files_root: str,
    search_subfolders: bool,
    fuzzy_match: bool,
    image_width_mode: str,          # fit_bubble | fixed
    image_max_width_cm: float,
    compress_images: bool,
    jpeg_quality: int,
    max_image_px: int,
    keep_media_cache: bool,
    video_thumbs: bool,
    video_thumb_width_cm: float,
    status_col: str | None,
    status_map: dict,
    orientation: str,  # portrait | landscape
    page_title: str,
    intro_text: str,
    page_numbers: bool,
    font_name: str = "Calibri",
    font_color_hex: str = "#000000",
    sender_font_size: int = 9,
    msg_font_size: int = 10,
    meta_font_size: int = 8,
):
    doc = Document()

    section = doc.sections[0]
    if orientation == "landscape":
        section.orientation = WD_ORIENTATION.LANDSCAPE
        new_width, new_height = section.page_height, section.page_width
        section.page_width = new_width
        section.page_height = new_height
    else:
        section.orientation = WD_ORIENTATION.PORTRAIT

    # Marges
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(0.7)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)

    if page_numbers:
        add_page_number(section)

    # Title + intro
    if page_title.strip():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text_with_linebreaks(
            p,
            page_title.strip(),
            font_size_pt=16,
            bold=True,
            font_name=font_name,
            font_color_hex=font_color_hex,
        )
        p.paragraph_format.space_after = Pt(8)

    if intro_text.strip():
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_text_with_linebreaks(
            p2,
            intro_text.strip(),
            font_size_pt=msg_font_size,
            bold=False,
            font_name=font_name,
            font_color_hex=font_color_hex,
        )
        p2.paragraph_format.space_after = Pt(10)

    out_dir = Path(out_path).parent
    cache_dir = out_dir / "_export_media_cache" if keep_media_cache else Path(tempfile.gettempdir()) / "chat_export_media_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)

    # LibreOffice : version V2.
    # On évite le grand tableau à 3 colonnes car LibreOffice affiche souvent
    # ses limites de cellules et donne un rendu "grille". Chaque message est
    # donc exporté comme un petit tableau indépendant, aligné à gauche ou à droite.

    # Validate columns
    if col_sender not in df.columns:
        raise ValueError(f"Colonne expéditeur introuvable: {col_sender}")
    if col_message not in df.columns:
        raise ValueError(f"Colonne message introuvable: {col_message}")
    if col_direction and col_direction not in df.columns:
        col_direction = None
    if media_path_col and media_path_col not in df.columns:
        media_path_col = None
    if status_col and status_col not in df.columns:
        status_col = None
    if owner_col and owner_col not in df.columns:
        owner_col = None

    # Compute usable width
    usable_width_cm = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
    usable_width_cm = max(10.0, usable_width_cm)

    gutter_cm = float(gutter_cm)
    gutter_cm = max(0.2, min(gutter_cm, 3.0))

    bubble_width_cm = max(4.0, min(float(bubble_width_cm), usable_width_cm - gutter_cm - 2.0))
    other_width_cm = max(2.0, usable_width_cm - bubble_width_cm - gutter_cm)

    bubble_spacing_pt = float(bubble_spacing_pt)
    bubble_spacing_pt = max(0.0, min(bubble_spacing_pt, 40.0))

    for _, r in df.iterrows():
        raw_sender_val = r.get(col_sender, "")
        sender_raw = _normalize_participant(raw_sender_val)

        # display name
        if sender_raw in display_name_map:
            sender = str(display_name_map.get(sender_raw, "")).strip() or empty_sender_display
        else:
            sender = sender_raw if sender_raw else empty_sender_display

        msg_val = r.get(col_message, "")
        msg = "" if pd.isna(msg_val) else clean_multiline_text(str(msg_val))

        # Optional: append extra mapped columns
        if extra_cols:
            extras = []
            for c in extra_cols:
                if c and (c in df.columns):
                    v = r.get(c, "")
                    if v is not None and not (isinstance(v, float) and pd.isna(v)):
                        s = str(v).strip()
                        if s:
                            extras.append(s)
            if extras:
                msg = (msg.rstrip() + "\n" if msg.strip() else "") + "\n".join(extras)

        timestamp = build_timestamp(r, ts_single_col, ts_date_col, ts_time_col)
        status_label = resolve_status_label(r, status_col, status_map) if status_col else ""

        # Determine side
        is_right = False
        if mode_side == "direction" and col_direction:
            d = normalize_direction(r.get(col_direction, ""))
            is_right = (d == "outgoing")
        elif mode_side == "sender" and right_participant_value is not None:
            if owner_col:
                owner_raw = _normalize_participant(r.get(owner_col, ""))
                is_right = (owner_raw == right_participant_value)
            else:
                is_right = (sender_raw == right_participant_value)

        # Media values
        media_values: list[str] = []
        if media_path_col:
            v = r.get(media_path_col, None)
            if v is not None and not (isinstance(v, float) and pd.isna(v)):
                media_values = split_media_values(str(v))

        # Skip fully empty (no msg + no media)
        if (not msg.strip()) and (not media_values):
            continue

        # Table indépendante pour une seule bulle.
        # C'est le rendu le plus fiable dans LibreOffice Writer : pas de grande grille,
        # pas de cellules vides visibles, pas de décalage de colonnes.
        bubble_table = doc.add_table(rows=1, cols=1)
        bubble_table.alignment = WD_TABLE_ALIGNMENT.RIGHT if is_right else WD_TABLE_ALIGNMENT.LEFT
        set_table_borders_none(bubble_table)
        set_table_fixed_layout(bubble_table)
        set_table_width_cm(bubble_table, bubble_width_cm)
        # Important LibreOffice/Word : éviter qu'une bulle soit coupée sur deux pages.
        # Si la bulle tient sur une page, Writer la bascule entière sur la page suivante.
        try:
            set_row_cant_split(bubble_table.rows[0])
        except Exception:
            pass

        bubble_cell = bubble_table.cell(0, 0)
        set_cell_border_none(bubble_cell)
        set_cell_width_cm(bubble_cell, bubble_width_cm)
        bubble_cell.width = Cm(bubble_width_cm)
        set_cell_margins(bubble_cell, top=90, start=160, bottom=90, end=160)

        # Couleur de bulle pilotée par le participant, sinon fallback gauche/droite
        bubble_color = participant_color_map.get(sender_raw) if participant_color_map else None
        if not bubble_color:
            bubble_color = default_right_color if is_right else default_left_color
        set_cell_shading(bubble_cell, bubble_color)

        # Sender line
        p_sender = bubble_cell.paragraphs[0]
        compact_paragraph(p_sender)
        p_sender.alignment = WD_ALIGN_PARAGRAPH.LEFT
        keep_paragraph_together(p_sender)
        add_text_with_linebreaks(
            p_sender,
            sender,
            font_size_pt=sender_font_size,
            bold=True,
            font_name=font_name,
            font_color_hex=font_color_hex,
        )

        # Body
        if msg.strip():
            p_body = bubble_cell.add_paragraph()
            compact_paragraph(p_body)
            p_body.alignment = WD_ALIGN_PARAGRAPH.LEFT
            keep_paragraph_together(p_body)
            add_text_with_linebreaks(
                p_body,
                msg,
                font_size_pt=msg_font_size,
                font_name=font_name,
                font_color_hex=font_color_hex,
            )

        # Medias
        for media_name in media_values:
            resolved = (
                resolve_media_path(media_name, files_root, recursive=search_subfolders, fuzzy=fuzzy_match)
                or resolve_media_path(media_name, str(out_dir), recursive=True, fuzzy=True)
            )

            p_media = bubble_cell.add_paragraph()
            compact_paragraph(p_media)
            p_media.alignment = WD_ALIGN_PARAGRAPH.LEFT
            keep_paragraph_together(p_media)

            if not resolved:
                add_text_with_linebreaks(
                    p_media,
                    f"[Pièce jointe] {media_name} — FICHIER NON RÉSOLU",
                    font_size_pt=meta_font_size,
                    font_name=font_name,
                )
                continue

            if is_image_file(resolved):
                insert_path = resolved
                if compress_images:
                    insert_path = compress_image_for_docx(insert_path, max_image_px, jpeg_quality, str(cache_dir))

                target_cm = float(image_max_width_cm)
                if image_width_mode == "fit_bubble":
                    target_cm = min(target_cm, bubble_width_cm)

                try:
                    run = p_media.add_run()
                    run.add_picture(insert_path, width=Cm(max(2.0, target_cm)))
                except Exception:
                    add_text_with_linebreaks(
                        p_media,
                        f"[Image] {os.path.basename(resolved)} — insertion impossible",
                        font_size_pt=meta_font_size,
                        font_name=font_name,
                    )

                p_name = bubble_cell.add_paragraph()
                compact_paragraph(p_name)
                keep_paragraph_together(p_name)
                add_text_with_linebreaks(
                    p_name,
                    os.path.basename(resolved),
                    font_size_pt=meta_font_size,
                    font_name=font_name,
                    font_color_hex=font_color_hex,
                )

            elif is_video_file(resolved):
                if video_thumbs:
                    thumb = make_video_thumbnail(resolved, str(cache_dir))
                    if thumb and os.path.exists(thumb):
                        target_cm = float(video_thumb_width_cm)
                        if image_width_mode == "fit_bubble":
                            target_cm = min(target_cm, bubble_width_cm)
                        try:
                            run = p_media.add_run()
                            run.add_picture(thumb, width=Cm(max(2.0, target_cm)))
                        except Exception:
                            pass

                add_text_with_linebreaks(
                    p_media,
                    f"[Vidéo] {os.path.basename(resolved)}",                    font_size_pt=meta_font_size,
                    font_name=font_name,
                )

            else:
                # Fichier non compatible (ni image ni vidéo)
                name_only = os.path.basename(resolved)

                run1 = p_media.add_run(name_only)
                run1.bold = True
                run1.underline = True
                run1.font.size = Pt(meta_font_size)
                run1.font.name = font_name
                try:
                    run1.font.color.rgb = RGBColor.from_string(_hex6(font_color_hex))
                except Exception:
                    pass

                # Badge simple (plus stable qu'une table dans une cellule)
                p_badge = bubble_cell.add_paragraph()
                compact_paragraph(p_badge)
                keep_paragraph_together(p_badge)
                run_badge = p_badge.add_run("NON COMPATIBLE")
                run_badge.bold = True
                run_badge.font.size = Pt(meta_font_size)
                run_badge.font.name = font_name
                try:
                    run_badge.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
                except Exception:
                    pass

        # status + timestamp always shown (default)
        if status_label or timestamp.strip():
            p_line = bubble_cell.add_paragraph()
            compact_paragraph(p_line)
            keep_paragraph_together(p_line)
            p_line.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_line.paragraph_format.tab_stops.add_tab_stop(
                Cm(max(1.0, bubble_width_cm - 0.3)),
                alignment=WD_TAB_ALIGNMENT.RIGHT,
                leader=WD_TAB_LEADER.SPACES,
            )
            if status_label:
                add_text_with_linebreaks(
                    p_line,
                    status_label,
                    font_size_pt=meta_font_size,
                    italic=True,
                    font_name=font_name,
                    font_color_hex=font_color_hex,
                )
            if timestamp.strip():
                p_line.add_run("\t")
                run_t = p_line.add_run(timestamp)
                run_t.font.size = Pt(meta_font_size)
                run_t.font.name = font_name
                try:
                    run_t.font.color.rgb = RGBColor.from_string(_hex6(font_color_hex))
                except Exception:
                    pass

        # Espacement vertical entre deux bulles.
        try:
            bubble_cell.paragraphs[-1].paragraph_format.space_after = Pt(0)
        except Exception:
            pass
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(bubble_spacing_pt)
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.line_spacing = 1

    doc.save(out_path)


# -----------------------------
# GUI
# -----------------------------
class App(ctk.CTk):
    def __init__(self):
        super().__init__()
        ctk.set_appearance_mode("Dark")
        ctk.set_default_color_theme("blue")

        self.title("Export conversation Excel -> LibreOffice Writer (bulles V4) — Dark Pro")
        self.geometry("1500x880")
        self.minsize(1200, 720)

        self.excel_path: str | None = None
        self.df: pd.DataFrame | None = None
        self.sheets: list[str] = []

        # Default bubble colors (fallback when no participant color)
        self.left_color = "#3D7DFF"
        self.right_color = "#7ED957"

        # Variables
        self.sheet_var = tk.StringVar(value="")
        self.col_sender_var = tk.StringVar(value="")
        self.col_msg_var = tk.StringVar(value="")
        self.col_dir_var = tk.StringVar(value="")

        self.ts_mode_var = tk.StringVar(value="split")
        self.ts_single_var = tk.StringVar(value="")
        self.ts_date_var = tk.StringVar(value="")
        self.ts_time_var = tk.StringVar(value="")

        self.side_mode_var = tk.StringVar(value="direction")
        self.owner_col_var = tk.StringVar(value="")
        self.right_participant_var = tk.StringVar(value="")
        self._participant_display_to_value: dict[str, str] = {}

        # Participant mapping
        self.display_name_map: dict[str, str] = {}      # source -> display
        self.participant_color_map: dict[str, str] = {} # source -> hex color
        self.empty_sender_name_var = tk.StringVar(value="Proprietaire")

        # Status
        self.show_status_var = tk.BooleanVar(value=False)
        self.status_col_var = tk.StringVar(value="")
        self.status_map = {"Yes": "Message effacé", "Trash": "Message à la corbeille"}

        # Preview
        self.filter_var = tk.StringVar(value="")
        self.preview_rows_var = tk.IntVar(value=50)

        # Extra columns
        self.extra_cols_selected: list[str] = []

        # Format tab
        self.orientation_var = tk.StringVar(value="portrait")
        self.font_name_var = tk.StringVar(value="Calibri")
        self.font_color_var = tk.StringVar(value="#000000")
        self.sender_font_size_var = tk.IntVar(value=9)
        self.msg_font_size_var = tk.IntVar(value=10)
        self.meta_font_size_var = tk.IntVar(value=8)
        self.bubble_width_var = tk.DoubleVar(value=10.5)
        self.gutter_cm_var = tk.DoubleVar(value=1.0)
        self.bubble_spacing_var = tk.DoubleVar(value=8.0)

        # Page options
        self.page_title_var = tk.StringVar(value="")
        self.intro_text_var = tk.StringVar(value="")
        self.page_numbers_var = tk.BooleanVar(value=False)

        # Options
        self.dedupe_var = tk.BooleanVar(value=True)
        self.open_word_var = tk.BooleanVar(value=True)

        # Files (beta) — medias
        self.media_col_var = tk.StringVar(value="")
        self.files_root_var = tk.StringVar(value="")
        self.search_subfolders_var = tk.BooleanVar(value=True)
        self.fuzzy_match_var = tk.BooleanVar(value=True)

        self.image_width_mode_var = tk.StringVar(value="fit_bubble")  # fit_bubble|fixed
        self.image_max_width_cm_var = tk.DoubleVar(value=8.5)

        self.compress_images_var = tk.BooleanVar(value=False)
        self.jpeg_quality_var = tk.IntVar(value=80)
        self.max_image_px_var = tk.IntVar(value=1920)
        self.keep_media_cache_var = tk.BooleanVar(value=True)

        self.video_thumb_var = tk.BooleanVar(value=True)
        self.video_thumb_width_cm_var = tk.DoubleVar(value=8.5)

        self.apply_ttk_dark_theme()
        self._build_ui()
        self.load_settings(silent=True)

        # Debounced sheet change
        self._sheet_after = None

        def _sheet_changed(*_):
            if not self.excel_path:
                return
            if self._sheet_after is not None:
                try:
                    self.after_cancel(self._sheet_after)
                except Exception:
                    pass
            self._sheet_after = self.after(80, self.load_sheet)

        try:
            self.sheet_var.trace_add("write", _sheet_changed)
        except Exception:
            pass

        # Enregistrer automatiquement les réglages à la fermeture.
        self.protocol("WM_DELETE_WINDOW", self.on_close)

    # ---------------------------
    # ttk theme dark
    # ---------------------------
    def apply_ttk_dark_theme(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass

        bg = "#1f1f1f"
        panel = "#2b2b2b"
        tab = "#2b2b2b"
        tab_sel = "#3a3a3a"
        fg = "#E6E6E6"
        border = "#3a3a3a"

        style.configure("TNotebook", background=bg, borderwidth=0)
        style.configure("TNotebook.Tab", background=tab, foreground=fg, padding=(12, 6), borderwidth=0)
        style.map(
            "TNotebook.Tab",
            background=[("selected", tab_sel), ("active", "#333333")],
            foreground=[("selected", "#FFFFFF"), ("active", "#FFFFFF")],
        )

        style.configure(
            "Treeview",
            background="#111111",
            fieldbackground="#111111",
            foreground=fg,
            bordercolor=border,
            borderwidth=0,
            rowheight=22,
        )
        style.map("Treeview", background=[("selected", "#2A5D9F")], foreground=[("selected", "#FFFFFF")])
        style.configure("Treeview.Heading", background=panel, foreground="#FFFFFF", relief="flat")
        style.map("Treeview.Heading", background=[("active", "#3a3a3a")], foreground=[("active", "#FFFFFF")])

        style.configure("TFrame", background=bg)

    # ---------------------------
    # UI helpers
    # ---------------------------
    def _field_row(self, parent, r, label, widget):
        ctk.CTkLabel(parent, text=label, text_color="#D0D0D0").grid(
            row=r, column=0, sticky="w", padx=(12, 10), pady=6
        )
        widget.grid(row=r, column=1, sticky="ew", padx=(0, 12), pady=6)

    def on_sheet_change(self, choice):
        """Callback pour changement de feuille: force la variable puis recharge."""
        try:
            self.sheet_var.set(choice)
        except Exception:
            pass
        self.refresh_current_sheet()

    def refresh_current_sheet(self):
        """Rafraîchit la feuille courante (relecture Excel + mise à jour UI)."""
        if not self.excel_path:
            return

        try:
            current = str(self.sheet_menu.get()).strip()
            if current:
                self.sheet_var.set(current)
        except Exception:
            current = self.sheet_var.get()

        try:
            self.sheets = list_sheets(self.excel_path)
            if self.sheets:
                self.sheet_menu.configure(values=self.sheets)
                if current not in self.sheets:
                    self.sheet_var.set(self.sheets[0])
            else:
                self.sheet_menu.configure(values=[""])
                self.sheet_var.set("")
        except Exception:
            pass

        self.load_sheet()

    # ---------------------------
    # UI build
    # ---------------------------
    def _build_ui(self):
        # Top bar
        top = ctk.CTkFrame(self, corner_radius=14)
        top.pack(fill="x", padx=12, pady=12)

        ctk.CTkButton(top, text="Ouvrir Excel…", width=140, command=self.open_excel).pack(
            side="left", padx=12, pady=10
        )
        ctk.CTkLabel(top, text="Feuille").pack(side="left", padx=(12, 8))

        self.sheet_menu = ctk.CTkOptionMenu(top, variable=self.sheet_var, values=[""], command=self.on_sheet_change)
        self.sheet_menu.pack(side="left", pady=10)

        ctk.CTkButton(
            top, text="Actualiser", width=110, fg_color="transparent", border_width=1, command=self.refresh_current_sheet
        ).pack(side="left", padx=(10, 0), pady=10)

        self.lbl_file = ctk.CTkLabel(top, text="Aucun fichier chargé", text_color="#A8A8A8")
        self.lbl_file.pack(side="left", padx=14)

        # Main split
        main = ctk.CTkFrame(self, corner_radius=14)
        main.pack(fill="both", expand=True, padx=12, pady=(0, 12))

        # Left panel
        left = ctk.CTkFrame(main, corner_radius=14, width=540)
        left.pack(side="left", fill="y", padx=(12, 8), pady=12)
        left.pack_propagate(False)

        left.grid_columnconfigure(0, weight=1)
        left.grid_rowconfigure(1, weight=1)

        ctk.CTkLabel(left, text="Paramétrage", font=ctk.CTkFont(size=16, weight="bold")).grid(
            row=0, column=0, sticky="w", padx=14, pady=(14, 8)
        )

        # Scroll host (tabs)
        scroll_host = ctk.CTkFrame(left, fg_color="transparent")
        scroll_host.grid(row=1, column=0, sticky="nsew", padx=12, pady=(0, 8))
        scroll_host.grid_columnconfigure(0, weight=1)
        scroll_host.grid_rowconfigure(0, weight=1)

        canvas = tk.Canvas(scroll_host, highlightthickness=0, bd=0, background="#1f1f1f")
        vscroll = ttk.Scrollbar(scroll_host, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=vscroll.set)

        canvas.grid(row=0, column=0, sticky="nsew")
        vscroll.grid(row=0, column=1, sticky="ns")

        nb_holder = ctk.CTkFrame(canvas, fg_color="transparent")
        win_id = canvas.create_window((0, 0), window=nb_holder, anchor="nw")

        def _on_nb_configure(_event=None):
            canvas.configure(scrollregion=canvas.bbox("all"))

        def _on_canvas_configure(event):
            canvas.itemconfigure(win_id, width=event.width)

        nb_holder.bind("<Configure>", _on_nb_configure)
        canvas.bind("<Configure>", _on_canvas_configure)

        def _on_mousewheel(event):
            if event.delta:
                canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        def _bind_wheel(_event=None):
            canvas.bind_all("<MouseWheel>", _on_mousewheel)

        def _unbind_wheel(_event=None):
            canvas.unbind_all("<MouseWheel>")

        canvas.bind("<Enter>", _bind_wheel)
        canvas.bind("<Leave>", _unbind_wheel)

        nb = ttk.Notebook(nb_holder)
        nb.pack(fill="both", expand=True)

        tab_base = ttk.Frame(nb)
        tab_format = ttk.Frame(nb)
        tab_options = ttk.Frame(nb)
        tab_files = ttk.Frame(nb)

        nb.add(tab_base, text="Base")
        nb.add(tab_format, text="Format")
        nb.add(tab_options, text="Options")
        nb.add(tab_files, text="Fichier (bêta)")

        # ---------------------------
        # TAB BASE
        # ---------------------------
        base = ctk.CTkFrame(tab_base, fg_color="transparent")
        base.pack(fill="both", expand=True)

        form = ctk.CTkFrame(base, fg_color="transparent")
        form.pack(fill="x", pady=(8, 0))
        form.grid_columnconfigure(1, weight=1)

        self.sender_menu = ctk.CTkOptionMenu(form, variable=self.col_sender_var, values=[""])
        self.msg_menu = ctk.CTkOptionMenu(form, variable=self.col_msg_var, values=[""])
        self.dir_menu = ctk.CTkOptionMenu(form, variable=self.col_dir_var, values=[""])

        self._field_row(form, 0, "Colonne expéditeur", self.sender_menu)
        self._field_row(form, 1, "Colonne message", self.msg_menu)
        self._field_row(form, 2, "Colonne direction", self.dir_menu)

        # Extra columns
        extra_box = ctk.CTkFrame(base, corner_radius=12)
        extra_box.pack(fill="x", pady=(8, 0))
        ctk.CTkLabel(extra_box, text="Infos supplémentaires", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 6)
        )

        extra_row = ctk.CTkFrame(extra_box, fg_color="transparent")
        extra_row.pack(fill="x", padx=12, pady=(0, 10))
        extra_row.grid_columnconfigure(1, weight=1)

        self.lbl_extra_cols = ctk.CTkLabel(extra_row, text="Aucune colonne additionnelle", text_color="#A8A8A8")
        self.lbl_extra_cols.grid(row=0, column=0, sticky="w")
        ctk.CTkButton(
            extra_row, text="Mapper…", width=110, fg_color="transparent", border_width=1,
            command=self.open_extra_columns_editor
        ).grid(row=0, column=1, sticky="e")

        # Timestamp
        ts_box = ctk.CTkFrame(base, corner_radius=12)
        ts_box.pack(fill="x", pady=(10, 6))
        ctk.CTkLabel(ts_box, text="Horodatage", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 4)
        )

        ts_row = ctk.CTkFrame(ts_box, fg_color="transparent")
        ts_row.pack(fill="x", padx=12)
        ctk.CTkRadioButton(
            ts_row, text="Date+Heure", variable=self.ts_mode_var, value="split", command=self._refresh_ts_mode
        ).pack(side="left", padx=(0, 12))
        ctk.CTkRadioButton(
            ts_row, text="1 colonne", variable=self.ts_mode_var, value="single", command=self._refresh_ts_mode
        ).pack(side="left")

        ts_form = ctk.CTkFrame(ts_box, fg_color="transparent")
        ts_form.pack(fill="x", padx=12, pady=(6, 10))
        ts_form.grid_columnconfigure(1, weight=1)

        self.ts_date_menu = ctk.CTkOptionMenu(ts_form, variable=self.ts_date_var, values=[""])
        self.ts_time_menu = ctk.CTkOptionMenu(ts_form, variable=self.ts_time_var, values=[""])
        self.ts_single_menu = ctk.CTkOptionMenu(ts_form, variable=self.ts_single_var, values=[""])

        self._field_row(ts_form, 0, "Colonne date", self.ts_date_menu)
        self._field_row(ts_form, 1, "Colonne heure", self.ts_time_menu)
        self._field_row(ts_form, 2, "Colonne timestamp", self.ts_single_menu)

        # Participants (rename / colors)
        names_box = ctk.CTkFrame(base, corner_radius=12)
        names_box.pack(fill="x", pady=6)
        ctk.CTkLabel(names_box, text="Noms des participants", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 6)
        )

        names_form = ctk.CTkFrame(names_box, fg_color="transparent")
        names_form.pack(fill="x", padx=12, pady=(0, 6))
        names_form.grid_columnconfigure(1, weight=1)
        self._field_row(names_form, 0, "Nom si participant vide", ctk.CTkEntry(names_form, textvariable=self.empty_sender_name_var))

        base_btns = ctk.CTkFrame(names_box, fg_color="transparent")
        base_btns.pack(fill="x", padx=12, pady=(0, 12))
        ctk.CTkButton(base_btns, text="Gérer participants…", fg_color="transparent", border_width=1,
                      command=self.open_participants_editor).pack(side="left")

        # Status
        status_box = ctk.CTkFrame(base, corner_radius=12)
        status_box.pack(fill="x", pady=(6, 0))
        ctk.CTkLabel(status_box, text="Statuts", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 6)
        )

        ctk.CTkCheckBox(
            status_box, text="Afficher le statut (effacé/corbeille…)", variable=self.show_status_var
        ).pack(anchor="w", padx=12, pady=(0, 6))

        status_form = ctk.CTkFrame(status_box, fg_color="transparent")
        status_form.pack(fill="x", padx=12, pady=(0, 6))
        status_form.grid_columnconfigure(1, weight=1)

        self.status_col_menu = ctk.CTkOptionMenu(status_form, variable=self.status_col_var, values=[""])
        self._field_row(status_form, 0, "Colonne statut", self.status_col_menu)

        self.btn_status_editor = ctk.CTkButton(
            status_box, text="Gérer statuts…", fg_color="transparent", border_width=1, command=self.open_status_editor
        )
        self.btn_status_editor.pack(anchor="w", padx=12, pady=(0, 12))

        def refresh_status_controls(*_):
            enabled = bool(self.show_status_var.get())
            self.status_col_menu.configure(state=("normal" if enabled else "disabled"))
            self.btn_status_editor.configure(state=("normal" if enabled else "disabled"))

        self.show_status_var.trace_add("write", refresh_status_controls)
        refresh_status_controls()

        # ---------------------------
        # TAB FORMAT
        # ---------------------------
        fmt = ctk.CTkFrame(tab_format, fg_color="transparent")
        fmt.pack(fill="both", expand=True)

        # Page box
        page_box = ctk.CTkFrame(fmt, corner_radius=12)
        page_box.pack(fill="x", padx=12, pady=(12, 8))
        ctk.CTkLabel(page_box, text="Page", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 6)
        )

        prow = ctk.CTkFrame(page_box, fg_color="transparent")
        prow.pack(fill="x", padx=12, pady=(0, 8))
        ctk.CTkLabel(prow, text="Orientation", text_color="#D0D0D0").pack(side="left")
        ctk.CTkOptionMenu(prow, values=["portrait", "landscape"], variable=self.orientation_var).pack(side="right")

        pgrid = ctk.CTkFrame(page_box, fg_color="transparent")
        pgrid.pack(fill="x", padx=12, pady=(0, 10))
        pgrid.grid_columnconfigure(1, weight=1)

        self._field_row(pgrid, 0, "Ajouter un titre", ctk.CTkEntry(pgrid, textvariable=self.page_title_var))
        ctk.CTkCheckBox(page_box, text="Numérotation (1 / 2)", variable=self.page_numbers_var).pack(
            anchor="w", padx=12, pady=(0, 10)
        )

        # Intro
        intro_box = ctk.CTkFrame(fmt, corner_radius=12)
        intro_box.pack(fill="x", padx=12, pady=(0, 8))
        ctk.CTkLabel(intro_box, text="Introduction", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 6)
        )
        intro_entry = ctk.CTkTextbox(intro_box, height=80)
        intro_entry.pack(fill="x", padx=12, pady=(0, 10))
        intro_entry.insert("1.0", "")
        self._intro_textbox = intro_entry

        # Bubbles
        style_box = ctk.CTkFrame(fmt, corner_radius=12)
        style_box.pack(fill="x", padx=12, pady=(0, 8))
        ctk.CTkLabel(style_box, text="Bulles", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 6)
        )

        bw_row = ctk.CTkFrame(style_box, fg_color="transparent")
        bw_row.pack(fill="x", padx=12, pady=(0, 0))
        ctk.CTkLabel(bw_row, text="Largeur bulles (cm)", text_color="#D0D0D0").pack(side="left")
        self.bw_value = ctk.CTkLabel(bw_row, text=f"{self.bubble_width_var.get():.1f}", text_color="#A8A8A8")
        self.bw_value.pack(side="right")

        def on_bw(v):
            self.bw_value.configure(text=f"{float(v):.1f}")

        ctk.CTkSlider(style_box, from_=4.0, to=22.0, number_of_steps=180, variable=self.bubble_width_var, command=on_bw).pack(
            fill="x", padx=12, pady=(2, 10)
        )

        gut_row = ctk.CTkFrame(style_box, fg_color="transparent")
        gut_row.pack(fill="x", padx=12, pady=(0, 0))
        ctk.CTkLabel(gut_row, text="Gouttière horizontale (cm)", text_color="#D0D0D0").pack(side="left")
        self.gut_value = ctk.CTkLabel(gut_row, text=f"{self.gutter_cm_var.get():.1f}", text_color="#A8A8A8")
        self.gut_value.pack(side="right")

        def on_gut(v):
            self.gut_value.configure(text=f"{float(v):.1f}")

        ctk.CTkSlider(style_box, from_=0.2, to=3.0, number_of_steps=140, variable=self.gutter_cm_var, command=on_gut).pack(
            fill="x", padx=12, pady=(2, 10)
        )

        sp_row = ctk.CTkFrame(style_box, fg_color="transparent")
        sp_row.pack(fill="x", padx=12, pady=(0, 0))
        ctk.CTkLabel(sp_row, text="Espacement vertical (pt)", text_color="#D0D0D0").pack(side="left")
        self.sp_value = ctk.CTkLabel(sp_row, text=f"{self.bubble_spacing_var.get():.0f}", text_color="#A8A8A8")
        self.sp_value.pack(side="right")

        def on_sp(v):
            self.sp_value.configure(text=f"{float(v):.0f}")

        ctk.CTkSlider(style_box, from_=0.0, to=30.0, number_of_steps=60, variable=self.bubble_spacing_var, command=on_sp).pack(
            fill="x", padx=12, pady=(2, 10)
        )

        # Typo
        typo_box = ctk.CTkFrame(fmt, corner_radius=12)
        typo_box.pack(fill="x", padx=12, pady=(0, 8))
        ctk.CTkLabel(typo_box, text="Typographie", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 6)
        )

        typo_form = ctk.CTkFrame(typo_box, fg_color="transparent")
        typo_form.pack(fill="x", padx=12, pady=(0, 10))
        for i in range(6):
            typo_form.grid_columnconfigure(i, weight=(1 if i in (1, 3, 5) else 0))

        common_fonts = ["Calibri", "Arial", "Times New Roman", "Segoe UI", "Verdana", "Courier New"]
        font_menu = ctk.CTkOptionMenu(typo_form, values=common_fonts, variable=self.font_name_var)

        ctk.CTkLabel(typo_form, text="Police", text_color="#D0D0D0").grid(row=0, column=0, sticky="w", padx=(0, 10), pady=6)
        font_menu.grid(row=0, column=1, sticky="ew", pady=6)

        ctk.CTkLabel(typo_form, text="Couleur", text_color="#D0D0D0").grid(row=0, column=2, sticky="w", padx=(18, 10), pady=6)

        def pick_font_color():
            c = colorchooser.askcolor(initialcolor=self.font_color_var.get(), title="Couleur du texte")
            if c and c[1]:
                self.font_color_var.set(c[1])  # pas d'appel apply_row ici (bug corrigé)

        color_btn = ctk.CTkButton(typo_form, text="Choisir…", width=110, fg_color="transparent", border_width=1, command=pick_font_color)
        color_btn.grid(row=0, column=3, sticky="w", pady=6)

        font_color_preview = ctk.CTkLabel(typo_form, textvariable=self.font_color_var, text_color="#A8A8A8")
        font_color_preview.grid(row=0, column=4, columnspan=2, sticky="e", pady=6)

        sizes = [str(i) for i in range(5, 15)]
        sender_menu = ctk.CTkOptionMenu(typo_form, values=sizes, command=lambda v: self.sender_font_size_var.set(int(v)))
        msg_menu = ctk.CTkOptionMenu(typo_form, values=sizes, command=lambda v: self.msg_font_size_var.set(int(v)))
        meta_menu = ctk.CTkOptionMenu(typo_form, values=sizes, command=lambda v: self.meta_font_size_var.set(int(v)))

        sender_menu.set(str(self.sender_font_size_var.get()))
        msg_menu.set(str(self.msg_font_size_var.get()))
        meta_menu.set(str(self.meta_font_size_var.get()))

        ctk.CTkLabel(typo_form, text="Taille exp.", text_color="#D0D0D0").grid(row=1, column=0, sticky="w", padx=(0, 10), pady=6)
        sender_menu.grid(row=1, column=1, sticky="ew", pady=6)

        ctk.CTkLabel(typo_form, text="Taille msg.", text_color="#D0D0D0").grid(row=1, column=2, sticky="w", padx=(18, 10), pady=6)
        msg_menu.grid(row=1, column=3, sticky="ew", pady=6)

        ctk.CTkLabel(typo_form, text="Taille meta", text_color="#D0D0D0").grid(row=1, column=4, sticky="w", padx=(18, 10), pady=6)
        meta_menu.grid(row=1, column=5, sticky="ew", pady=6)

        # ---------------------------
        # TAB OPTIONS
        # ---------------------------
        opt = ctk.CTkFrame(tab_options, fg_color="transparent")
        opt.pack(fill="both", expand=True)

        opt_box = ctk.CTkFrame(opt, corner_radius=12)
        opt_box.pack(fill="x", padx=12, pady=(12, 8))

        opt_form = ctk.CTkFrame(opt_box, fg_color="transparent")
        opt_form.pack(fill="x", padx=12, pady=(10, 10))
        opt_form.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(opt_form, text="Options export", font=ctk.CTkFont(weight="bold")).grid(
            row=0, column=0, columnspan=2, sticky="w", pady=(0, 6)
        )

        ctk.CTkCheckBox(opt_form, text="Supprimer doublons (From+Body+Date)", variable=self.dedupe_var).grid(
            row=1, column=0, columnspan=2, sticky="w", pady=(6, 2)
        )
        ctk.CTkCheckBox(opt_form, text="Ouvrir Word après génération", variable=self.open_word_var).grid(
            row=2, column=0, columnspan=2, sticky="w", pady=(6, 2)
        )

        settings_row = ctk.CTkFrame(opt_form, fg_color="transparent")
        settings_row.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(12, 0))
        ctk.CTkButton(
            settings_row, text="Enregistrer paramètres", fg_color="transparent", border_width=1,
            command=self.save_settings
        ).pack(side="left", padx=(0, 8))
        ctk.CTkButton(
            settings_row, text="Charger paramètres", fg_color="transparent", border_width=1,
            command=lambda: self.load_settings(silent=False)
        ).pack(side="left")

        # Owner bubble side mode
        owner_opt = ctk.CTkFrame(opt, corner_radius=12)
        owner_opt.pack(fill="x", padx=12, pady=(0, 8))
        ctk.CTkLabel(owner_opt, text="Bulle droite (owner)", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 4)
        )

        owner_row = ctk.CTkFrame(owner_opt, fg_color="transparent")
        owner_row.pack(fill="x", padx=12, pady=(0, 6))
        ctk.CTkRadioButton(owner_row, text="Via Direction", variable=self.side_mode_var, value="direction").pack(
            side="left", padx=(0, 12)
        )
        ctk.CTkRadioButton(owner_row, text="Via valeur d’une colonne", variable=self.side_mode_var, value="sender").pack(side="left")

        owner_grid = ctk.CTkFrame(owner_opt, fg_color="transparent")
        owner_grid.pack(fill="x", padx=12, pady=(0, 12))
        owner_grid.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(owner_grid, text="Colonne à tester", text_color="#D0D0D0").grid(
            row=0, column=0, sticky="w", padx=(0, 10), pady=6
        )
        self.owner_col_menu = ctk.CTkOptionMenu(
            owner_grid,
            variable=self.owner_col_var,
            values=[""],
            command=lambda _v: self.update_owner_values()
        )
        self.owner_col_menu.grid(row=0, column=1, sticky="ew", pady=6)

        ctk.CTkLabel(owner_grid, text="Valeur à mettre à droite", text_color="#D0D0D0").grid(
            row=1, column=0, sticky="w", padx=(0, 10), pady=6
        )
        self.part_menu = ctk.CTkOptionMenu(owner_grid, variable=self.right_participant_var, values=[""])
        self.part_menu.grid(row=1, column=1, sticky="ew", pady=6)

        # ---------------------------
        # TAB FICHIER (BETA)
        # ---------------------------
        files = ctk.CTkFrame(tab_files, fg_color="transparent")
        files.pack(fill="both", expand=True)

        box = ctk.CTkFrame(files, corner_radius=12)
        box.pack(fill="x", padx=12, pady=(12, 8))
        ctk.CTkLabel(box, text="Fichiers / pièces jointes", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 6)
        )

        row = ctk.CTkFrame(box, fg_color="transparent")
        row.pack(fill="x", padx=12, pady=(0, 10))
        row.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(row, text="Dossier fichiers", text_color="#D0D0D0").grid(row=0, column=0, sticky="w", padx=(0, 10), pady=6)
        ctk.CTkEntry(row, textvariable=self.files_root_var).grid(row=0, column=1, sticky="ew", pady=6)

        def browse_files_root():
            path = filedialog.askdirectory(title="Choisir le dossier contenant les fichiers")
            if path:
                self.files_root_var.set(path)

        ctk.CTkButton(row, text="Parcourir…", width=110, command=browse_files_root).grid(row=0, column=2, padx=(10, 0), pady=6)

        media_row = ctk.CTkFrame(box, fg_color="transparent")
        media_row.pack(fill="x", padx=12, pady=(0, 0))
        media_row.grid_columnconfigure(1, weight=1)
        self.media_menu = ctk.CTkOptionMenu(media_row, variable=self.media_col_var, values=[""])
        self._field_row(media_row, 0, "Colonne média (nom/chemin)", self.media_menu)

        ctk.CTkCheckBox(files, text="Rechercher dans les sous-dossiers", variable=self.search_subfolders_var).pack(
            anchor="w", padx=24, pady=(6, 0)
        )
        ctk.CTkCheckBox(files, text="Correspondance partielle des noms", variable=self.fuzzy_match_var).pack(
            anchor="w", padx=24, pady=(4, 0)
        )

        media_opts = ctk.CTkFrame(files, corner_radius=12)
        media_opts.pack(fill="x", padx=12, pady=(8, 8))
        ctk.CTkLabel(media_opts, text="Options médias", font=ctk.CTkFont(weight="bold")).pack(
            anchor="w", padx=12, pady=(10, 6)
        )

        mo = ctk.CTkFrame(media_opts, fg_color="transparent")
        mo.pack(fill="x", padx=12, pady=(0, 10))
        mo.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(mo, text="Taille image", text_color="#D0D0D0").grid(row=0, column=0, sticky="w", padx=(0, 10), pady=6)
        ctk.CTkOptionMenu(mo, values=["fit_bubble", "fixed"], variable=self.image_width_mode_var).grid(
            row=0, column=1, sticky="ew", pady=6
        )

        ctk.CTkLabel(mo, text="Largeur max image (cm)", text_color="#D0D0D0").grid(row=1, column=0, sticky="w", padx=(0, 10), pady=6)
        ctk.CTkSlider(mo, from_=2.0, to=15.0, number_of_steps=130, variable=self.image_max_width_cm_var).grid(
            row=1, column=1, sticky="ew", pady=6
        )

        ctk.CTkCheckBox(media_opts, text="Compresser les images (aperçu)", variable=self.compress_images_var).pack(
            anchor="w", padx=12, pady=(0, 6)
        )

        comp = ctk.CTkFrame(media_opts, fg_color="transparent")
        comp.pack(fill="x", padx=12, pady=(0, 8))
        comp.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(comp, text="Qualité JPEG", text_color="#D0D0D0").grid(row=0, column=0, sticky="w", padx=(0, 10), pady=6)
        ctk.CTkSlider(
            comp, from_=30, to=95, number_of_steps=65,
            variable=tk.DoubleVar(value=float(self.jpeg_quality_var.get())),
            command=lambda v: self.jpeg_quality_var.set(int(float(v))),
        ).grid(row=0, column=1, sticky="ew", pady=6)

        ctk.CTkLabel(comp, text="Max pixels (plus grand côté)", text_color="#D0D0D0").grid(row=1, column=0, sticky="w", padx=(0, 10), pady=6)
        px_menu = ctk.CTkOptionMenu(comp, values=["1280", "1920", "2560", "3840"], command=lambda v: self.max_image_px_var.set(int(v)))
        px_menu.set(str(self.max_image_px_var.get()))
        px_menu.grid(row=1, column=1, sticky="ew", pady=6)

        ctk.CTkCheckBox(media_opts, text="Conserver cache média (_export_media_cache)", variable=self.keep_media_cache_var).pack(
            anchor="w", padx=12, pady=(0, 10)
        )

        ctk.CTkCheckBox(media_opts, text="Miniature vidéo (ffmpeg)", variable=self.video_thumb_var).pack(
            anchor="w", padx=12, pady=(0, 6)
        )

        vt = ctk.CTkFrame(media_opts, fg_color="transparent")
        vt.pack(fill="x", padx=12, pady=(0, 10))
        vt.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(vt, text="Largeur miniature vidéo (cm)", text_color="#D0D0D0").grid(row=0, column=0, sticky="w", padx=(0, 10), pady=6)
        ctk.CTkSlider(vt, from_=2.0, to=15.0, number_of_steps=130, variable=self.video_thumb_width_cm_var).grid(
            row=0, column=1, sticky="ew", pady=6
        )

        ctk.CTkLabel(files, text="⚠ Fonction bêta – dépend de la structure des exports", text_color="#999999").pack(
            anchor="w", padx=24, pady=(0, 0)
        )

        # Bottom actions (always visible)
        actions = ctk.CTkFrame(left, fg_color="transparent")
        actions.grid(row=2, column=0, sticky="ew", padx=12, pady=(0, 12))

        settings_actions = ctk.CTkFrame(actions, fg_color="transparent")
        settings_actions.pack(fill="x", pady=(0, 8))
        settings_actions.grid_columnconfigure(0, weight=1)
        settings_actions.grid_columnconfigure(1, weight=1)

        ctk.CTkButton(
            settings_actions,
            text="Enregistrer paramètres",
            fg_color="transparent",
            border_width=1,
            command=self.save_settings,
        ).grid(row=0, column=0, sticky="ew", padx=(0, 4))

        ctk.CTkButton(
            settings_actions,
            text="Charger paramètres",
            fg_color="transparent",
            border_width=1,
            command=lambda: self.load_settings(silent=False),
        ).grid(row=0, column=1, sticky="ew", padx=(4, 0))

        self.export_btn = ctk.CTkButton(actions, text="Exporter DOCX LibreOffice V5…", command=self.export_word, state="disabled")
        self.export_btn.pack(fill="x")

        # ---------------------------
        # Right panel (preview)
        # ---------------------------
        right = ctk.CTkFrame(main, corner_radius=14)
        right.pack(side="right", fill="both", expand=True, padx=(8, 12), pady=12)

        header = ctk.CTkFrame(right, fg_color="transparent")
        header.pack(fill="x", padx=12, pady=(12, 6))
        ctk.CTkLabel(header, text="Aperçu", font=ctk.CTkFont(size=16, weight="bold")).pack(side="left")

        tools = ctk.CTkFrame(right, fg_color="transparent")
        tools.pack(fill="x", padx=12, pady=(0, 8))

        ctk.CTkLabel(tools, text="Filtre").pack(side="left", padx=(0, 8))
        self.filter_entry = ctk.CTkEntry(
            tools, textvariable=self.filter_var, width=360, placeholder_text="Filtrer l’aperçu (toutes colonnes, non exporté)"
        )
        self.filter_entry.pack(side="left")
        self.filter_entry.bind("<KeyRelease>", lambda e: self._render_preview())

        ctk.CTkLabel(tools, text="Lignes").pack(side="left", padx=(16, 8))
        self.rows_menu = ctk.CTkOptionMenu(tools, values=["20", "50", "100", "200", "500"], command=lambda v: self._set_preview_rows(v))
        self.rows_menu.set("50")
        self.rows_menu.pack(side="left")

        ctk.CTkButton(tools, text="Effacer", fg_color="transparent", border_width=1, command=self._clear_filter).pack(
            side="left", padx=12
        )

        holder = ctk.CTkFrame(right)
        holder.pack(fill="both", expand=True, padx=12, pady=(0, 12))

        self.tree = ttk.Treeview(holder, columns=(), show="headings")
        hsb = ttk.Scrollbar(holder, orient="horizontal", command=self.tree.xview)
        self.tree.configure(xscrollcommand=hsb.set)
        self.tree.pack(side="top", fill="both", expand=True)
        hsb.pack(side="bottom", fill="x")

        self._refresh_ts_mode()

    # ---------------------------
    # UI events
    # ---------------------------
    def _refresh_ts_mode(self):
        split = (self.ts_mode_var.get() == "split")
        if split:
            self.ts_date_menu.grid()
            self.ts_time_menu.grid()
            self.ts_single_menu.grid_remove()
        else:
            self.ts_date_menu.grid_remove()
            self.ts_time_menu.grid_remove()
            self.ts_single_menu.grid()

    def _set_preview_rows(self, v):
        try:
            self.preview_rows_var.set(int(v))
        except Exception:
            self.preview_rows_var.set(50)
        self._render_preview()

    def _clear_filter(self):
        self.filter_var.set("")
        self._render_preview()

    # ---------------------------
    # Settings save/load
    # ---------------------------
    def _get_textbox_text(self, textbox, fallback_var=None):
        try:
            return textbox.get("1.0", "end").strip()
        except Exception:
            return fallback_var.get().strip() if fallback_var is not None else ""

    def collect_settings(self) -> dict:
        return {
            "version": 3,
            "columns": {
                "sender": self.col_sender_var.get(),
                "message": self.col_msg_var.get(),
                "direction": self.col_dir_var.get(),
                "owner_col": self.owner_col_var.get(),
                "owner_value": self.right_participant_var.get(),
                "ts_mode": self.ts_mode_var.get(),
                "ts_single": self.ts_single_var.get(),
                "ts_date": self.ts_date_var.get(),
                "ts_time": self.ts_time_var.get(),
                "media": self.media_col_var.get(),
                "status": self.status_col_var.get(),
            },
            "format": {
                "orientation": self.orientation_var.get(),
                "font_name": self.font_name_var.get(),
                "font_color": self.font_color_var.get() or "#000000",                "sender_font_size": int(self.sender_font_size_var.get()),
                "msg_font_size": int(self.msg_font_size_var.get()),
                "meta_font_size": int(self.meta_font_size_var.get()),
                "bubble_width": float(self.bubble_width_var.get()),
                "gutter_cm": float(self.gutter_cm_var.get()),
                "bubble_spacing": float(self.bubble_spacing_var.get()),
                "left_color": self.left_color,
                "right_color": self.right_color,
            },
            "page": {
                "title": self.page_title_var.get(),
                "intro": self._get_textbox_text(getattr(self, "_intro_textbox", None), self.intro_text_var),
                "page_numbers": bool(self.page_numbers_var.get()),
            },
            "options": {
                "dedupe": bool(self.dedupe_var.get()),
                "open_word": bool(self.open_word_var.get()),
                "side_mode": self.side_mode_var.get(),
                "show_status": bool(self.show_status_var.get()),
                "empty_sender": self.empty_sender_name_var.get(),
            },
            "participants": {
                "display_name_map": self.display_name_map,
                "participant_color_map": self.participant_color_map,
            },
            "status_map": self.status_map,
            "extra_cols_selected": self.extra_cols_selected,
            "files": {
                "root": self.files_root_var.get(),
                "search_subfolders": bool(self.search_subfolders_var.get()),
                "fuzzy_match": bool(self.fuzzy_match_var.get()),
                "image_width_mode": self.image_width_mode_var.get(),
                "image_max_width_cm": float(self.image_max_width_cm_var.get()),
                "compress_images": bool(self.compress_images_var.get()),
                "jpeg_quality": int(self.jpeg_quality_var.get()),
                "max_image_px": int(self.max_image_px_var.get()),
                "keep_media_cache": bool(self.keep_media_cache_var.get()),
                "video_thumb": bool(self.video_thumb_var.get()),
                "video_thumb_width_cm": float(self.video_thumb_width_cm_var.get()),
            },
        }

    def apply_settings(self, data: dict):
        columns = data.get("columns", {})
        self.col_sender_var.set(columns.get("sender", self.col_sender_var.get()))
        self.col_msg_var.set(columns.get("message", self.col_msg_var.get()))
        self.col_dir_var.set(columns.get("direction", self.col_dir_var.get()))
        self.owner_col_var.set(columns.get("owner_col", self.owner_col_var.get()))
        self.right_participant_var.set(columns.get("owner_value", self.right_participant_var.get()))
        self.ts_mode_var.set(columns.get("ts_mode", self.ts_mode_var.get()))
        self.ts_single_var.set(columns.get("ts_single", self.ts_single_var.get()))
        self.ts_date_var.set(columns.get("ts_date", self.ts_date_var.get()))
        self.ts_time_var.set(columns.get("ts_time", self.ts_time_var.get()))
        self.media_col_var.set(columns.get("media", self.media_col_var.get()))
        self.status_col_var.set(columns.get("status", self.status_col_var.get()))

        fmt = data.get("format", {})
        self.orientation_var.set(fmt.get("orientation", self.orientation_var.get()))
        self.font_name_var.set(fmt.get("font_name", self.font_name_var.get()))
        self.font_color_var.set(fmt.get("font_color", "#000000") or "#000000")
        self.sender_font_size_var.set(int(fmt.get("sender_font_size", self.sender_font_size_var.get())))
        self.msg_font_size_var.set(int(fmt.get("msg_font_size", self.msg_font_size_var.get())))
        self.meta_font_size_var.set(int(fmt.get("meta_font_size", self.meta_font_size_var.get())))
        self.bubble_width_var.set(float(fmt.get("bubble_width", self.bubble_width_var.get())))
        self.gutter_cm_var.set(float(fmt.get("gutter_cm", self.gutter_cm_var.get())))
        self.bubble_spacing_var.set(float(fmt.get("bubble_spacing", self.bubble_spacing_var.get())))
        self.left_color = fmt.get("left_color", self.left_color)
        self.right_color = fmt.get("right_color", self.right_color)

        page = data.get("page", {})
        self.page_title_var.set(page.get("title", self.page_title_var.get()))
        self.page_numbers_var.set(bool(page.get("page_numbers", self.page_numbers_var.get())))
        if hasattr(self, "_intro_textbox"):
            try:
                self._intro_textbox.delete("1.0", "end")
                self._intro_textbox.insert("1.0", page.get("intro", ""))
            except Exception:
                pass

        opt = data.get("options", {})
        self.dedupe_var.set(bool(opt.get("dedupe", self.dedupe_var.get())))
        self.open_word_var.set(bool(opt.get("open_word", self.open_word_var.get())))
        self.side_mode_var.set(opt.get("side_mode", self.side_mode_var.get()))
        self.show_status_var.set(bool(opt.get("show_status", self.show_status_var.get())))
        self.empty_sender_name_var.set(opt.get("empty_sender", self.empty_sender_name_var.get()))

        participants = data.get("participants", {})
        self.display_name_map = dict(participants.get("display_name_map", self.display_name_map))
        self.participant_color_map = dict(participants.get("participant_color_map", self.participant_color_map))
        self.status_map = dict(data.get("status_map", self.status_map))
        self.extra_cols_selected = list(data.get("extra_cols_selected", self.extra_cols_selected))

        files = data.get("files", {})
        self.files_root_var.set(files.get("root", self.files_root_var.get()))
        self.search_subfolders_var.set(bool(files.get("search_subfolders", self.search_subfolders_var.get())))
        self.fuzzy_match_var.set(bool(files.get("fuzzy_match", self.fuzzy_match_var.get())))
        self.image_width_mode_var.set(files.get("image_width_mode", self.image_width_mode_var.get()))
        self.image_max_width_cm_var.set(float(files.get("image_max_width_cm", self.image_max_width_cm_var.get())))
        self.compress_images_var.set(bool(files.get("compress_images", self.compress_images_var.get())))
        self.jpeg_quality_var.set(int(files.get("jpeg_quality", self.jpeg_quality_var.get())))
        self.max_image_px_var.set(int(files.get("max_image_px", self.max_image_px_var.get())))
        self.keep_media_cache_var.set(bool(files.get("keep_media_cache", self.keep_media_cache_var.get())))
        self.video_thumb_var.set(bool(files.get("video_thumb", self.video_thumb_var.get())))
        self.video_thumb_width_cm_var.set(float(files.get("video_thumb_width_cm", self.video_thumb_width_cm_var.get())))

        self._refresh_ts_mode()
        self.update_format_value_labels()
        self._refresh_extra_cols_label()
        if self.df is not None:
            self.update_owner_values()


    def _refresh_extra_cols_label(self):
        try:
            selected = self.extra_cols_selected or []
            if selected:
                shown = ", ".join(selected[:3])
                if len(selected) > 3:
                    shown += f" (+{len(selected)-3})"
                self.lbl_extra_cols.configure(text=f"Colonnes: {shown}")
            else:
                self.lbl_extra_cols.configure(text="Aucune colonne additionnelle")
        except Exception:
            pass

    def update_format_value_labels(self):
        try:
            self.bw_value.configure(text=f"{float(self.bubble_width_var.get()):.1f}")
        except Exception:
            pass
        try:
            self.gut_value.configure(text=f"{float(self.gutter_cm_var.get()):.1f}")
        except Exception:
            pass
        try:
            self.sp_value.configure(text=f"{float(self.bubble_spacing_var.get()):.0f}")
        except Exception:
            pass

    def save_settings(self, silent: bool = False):
        try:
            SETTINGS_PATH.parent.mkdir(parents=True, exist_ok=True)
            SETTINGS_PATH.write_text(
                json.dumps(self.collect_settings(), ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            if not silent:
                messagebox.showinfo("Paramètres", f"Paramètres enregistrés :\n{SETTINGS_PATH}")
            return True
        except Exception as e:
            if not silent:
                messagebox.showerror("Paramètres", f"Impossible d’enregistrer les paramètres :\n{e}")
            return False

    def on_close(self):
        """Sauvegarde automatique des réglages puis fermeture de l'application."""
        try:
            self.save_settings(silent=True)
        finally:
            self.destroy()

    def load_settings(self, silent: bool = False):
        if not SETTINGS_PATH.exists():
            if not silent:
                messagebox.showinfo("Paramètres", "Aucun paramétrage enregistré pour l’instant.")
            return
        try:
            data = json.loads(SETTINGS_PATH.read_text(encoding="utf-8"))
            self.apply_settings(data)
            if not silent:
                messagebox.showinfo("Paramètres", "Paramètres chargés.")
        except Exception as e:
            if not silent:
                messagebox.showerror("Paramètres", f"Impossible de charger les paramètres :\n{e}")

    # ---------------------------
    # Excel load + preview
    # ---------------------------
    def open_excel(self):
        path = filedialog.askopenfilename(
            title="Choisir un fichier Excel",
            filetypes=[("Excel", "*.xlsx *.xls"), ("Tous fichiers", "*.*")]
        )
        if not path:
            return

        self.excel_path = path
        self.lbl_file.configure(text=os.path.basename(path))

        self.sheets = list_sheets(path)
        if self.sheets:
            self.sheet_menu.configure(values=self.sheets)
            self.sheet_var.set(self.sheets[0])
        else:
            self.sheet_menu.configure(values=[""])
            self.sheet_var.set("")
        self.load_sheet()

    def load_sheet(self):
        if not self.excel_path:
            return

        sheet = self.sheet_var.get() if self.sheet_var.get() else None
        try:
            self.df = load_excel_smart(self.excel_path, sheet_name=sheet)
        except Exception as e:
            messagebox.showerror("Erreur", f"Impossible de lire l'Excel:\n{e}")
            return

        cols = list(self.df.columns)

        # Preserve selections
        prev = {
            "sender": self.col_sender_var.get(),
            "msg": self.col_msg_var.get(),
            "dir": self.col_dir_var.get(),
            "ts_date": self.ts_date_var.get(),
            "ts_time": self.ts_time_var.get(),
            "ts_single": self.ts_single_var.get(),
            "media": self.media_col_var.get(),
            "status": self.status_col_var.get(),
            "owner_col": self.owner_col_var.get(),
        }
        values = [""] + cols

        def pick_first(pref):
            for p in pref:
                if p in cols:
                    return p
            return cols[0] if cols else ""

        # 1) configure menus only (bug corrigé : pas de .set() ici)
        for menu in (
            self.sender_menu, self.msg_menu, self.dir_menu,
            self.ts_date_menu, self.ts_time_menu, self.ts_single_menu,
            self.media_menu, self.status_col_menu, self.owner_col_menu
        ):
            menu.configure(values=values)

        def keep_or_default(cur, default):
            return cur if (cur and cur in cols) else default

        # 2) set variables once
        self.col_sender_var.set(keep_or_default(prev["sender"], pick_first(["From", "Sender", "Author"])))
        self.col_msg_var.set(keep_or_default(prev["msg"], pick_first(["Body", "Transcript", "Message", "Text"])))
        self.col_dir_var.set(keep_or_default(prev["dir"], pick_first(["Direction", "dir", "Type"])))
        self.ts_date_var.set(keep_or_default(prev["ts_date"], pick_first(["Timestamp-Date", "Date", "timestamp_date"])))
        self.ts_time_var.set(keep_or_default(prev["ts_time"], pick_first(["Timestamp-Time", "Time", "timestamp_time"])))
        self.ts_single_var.set(keep_or_default(prev["ts_single"], pick_first(["Timestamp", "DateTime", "Created At"])))
        self.media_col_var.set(keep_or_default(prev["media"], pick_first(["Attachment #1", "MediaPath", "AttachmentPath", "Attachment", "Media"])))
        self.owner_col_var.set(keep_or_default(prev.get("owner_col", ""), self.col_sender_var.get()))

        if prev.get("status") and prev["status"] in cols:
            self.status_col_var.set(prev["status"])
        elif "Deleted" in cols and not self.status_col_var.get().strip():
            self.status_col_var.set("Deleted")

        # Participants list
        sender_col = self.col_sender_var.get()
        participants: list[str] = []
        if sender_col and sender_col in cols:
            raw_values = self.df[sender_col].fillna("").astype(str).unique()
            for v in raw_values:
                participants.append(_normalize_participant(v))
            participants = sorted(set(participants))

        self.update_owner_values()

        # Defaults mapping
        for p in participants:
            self.display_name_map.setdefault(p, p)
        self.display_name_map.setdefault("", self.empty_sender_name_var.get().strip() or "Unknown")

        # Default files root to Excel directory
        if not self.files_root_var.get().strip() and self.excel_path:
            self.files_root_var.set(str(Path(self.excel_path).parent))

        self.export_btn.configure(state="normal")
        self._render_preview()

    def update_owner_values(self):
        """Met à jour les valeurs disponibles pour placer les bulles à droite.

        Cette liste dépend de la colonne choisie dans "Colonne à tester".
        Cela évite de forcer l'utilisation de la première colonne ou de la colonne expéditeur.
        """
        if self.df is None or self.df.empty:
            self._participant_display_to_value = {}
            self.part_menu.configure(values=[""])
            self.right_participant_var.set("")
            return

        owner_col = self.owner_col_var.get().strip()
        if not owner_col or owner_col not in self.df.columns:
            self._participant_display_to_value = {}
            self.part_menu.configure(values=[""])
            self.right_participant_var.set("")
            return

        values: list[str] = []
        seen: set[str] = set()
        for v in self.df[owner_col].tolist():
            norm = _normalize_participant(v)
            if norm not in seen:
                seen.add(norm)
                values.append(norm)

        values = sorted(values, key=lambda x: (x == "", x.lower()))

        display_values: list[str] = []
        self._participant_display_to_value = {}
        for raw in values:
            label = "(valeur vide)" if raw == "" else raw
            display_values.append(label)
            self._participant_display_to_value[label] = raw

        self.part_menu.configure(values=(display_values if display_values else [""]))
        current_label = self.right_participant_var.get()
        if display_values:
            if current_label not in display_values:
                self.right_participant_var.set(display_values[0])
        else:
            self.right_participant_var.set("")

    def _render_preview(self):
        self.tree.delete(*self.tree.get_children())
        if self.df is None or self.df.empty:
            self.tree["columns"] = ()
            return

        dfv = self.df
        filt = self.filter_var.get().strip().lower()
        if filt:
            mask = pd.Series(False, index=dfv.index)
            needle = re.escape(filt)
            for c in dfv.columns:
                s = dfv[c].astype(str).str.lower()
                mask = mask | s.str.contains(needle, na=False)
            dfv = dfv[mask]

        n = max(1, int(self.preview_rows_var.get()))
        preview = dfv.head(n)
        cols = list(preview.columns)

        self.tree["columns"] = cols
        for c in cols:
            self.tree.heading(c, text=c)
            self.tree.column(c, width=140, stretch=False)

        for _, row in preview.iterrows():
            vals = []
            for c in cols:
                v = row.get(c, "")
                if pd.isna(v):
                    v = ""
                s = str(v).replace("\r\n", "\n").replace("\r", "\n")
                s_one = re.sub(r"\s+", " ", s).strip()
                if len(s_one) > 180:
                    s_one = s_one[:180] + "…"
                vals.append(s_one)
            self.tree.insert("", "end", values=vals)

    # ---------------------------
    # Extra columns editor
    # ---------------------------
    def open_extra_columns_editor(self):
        if self.df is None or self.df.empty:
            messagebox.showinfo("Info", "Chargez d'abord un Excel.")
            return

        win = ctk.CTkToplevel(self)
        win.title("Colonnes additionnelles")
        win.geometry("520x520")
        win.grab_set()

        ctk.CTkLabel(
            win, text="Choisir des colonnes à ajouter au contenu du message",
            font=ctk.CTkFont(size=14, weight="bold")
        ).pack(anchor="w", padx=12, pady=(12, 8))

        sf = ctk.CTkScrollableFrame(win, corner_radius=12)
        sf.pack(fill="both", expand=True, padx=12, pady=(0, 12))

        cols = list(self.df.columns)
        vars_by_col: dict[str, tk.BooleanVar] = {}
        for c in cols:
            v = tk.BooleanVar(value=(c in self.extra_cols_selected))
            vars_by_col[c] = v
            ctk.CTkCheckBox(sf, text=c, variable=v).pack(anchor="w", padx=12, pady=4)

        btns = ctk.CTkFrame(win, fg_color="transparent")
        btns.pack(fill="x", padx=12, pady=(0, 12))

        def apply_and_close():
            selected = [c for c, v in vars_by_col.items() if v.get()]
            self.extra_cols_selected = selected
            if selected:
                shown = ", ".join(selected[:3])
                if len(selected) > 3:
                    shown += f" (+{len(selected)-3})"
                self.lbl_extra_cols.configure(text=f"Colonnes: {shown}")
            else:
                self.lbl_extra_cols.configure(text="Aucune colonne additionnelle")
            win.destroy()

        ctk.CTkButton(btns, text="OK", width=120, command=apply_and_close).pack(side="right")
        ctk.CTkButton(btns, text="Fermer", width=120, fg_color="transparent", border_width=1, command=win.destroy).pack(
            side="right", padx=(0, 10)
        )

        win.bind("<Return>", lambda _e: apply_and_close())

    # ---------------------------
    # Participants editor
    # ---------------------------
    def open_participants_editor(self):
        if self.df is None or self.df.empty:
            messagebox.showinfo("Info", "Chargez d'abord un Excel.")
            return

        win = ctk.CTkToplevel(self)
        win.title("Participants — renommage")
        win.geometry("780x520")
        win.minsize(720, 480)
        win.grab_set()

        ctk.CTkLabel(
            win, text="Renommage participants (source → affiché)", font=ctk.CTkFont(size=15, weight="bold")
        ).pack(anchor="w", padx=12, pady=(12, 8))

        holder = ctk.CTkFrame(win)
        holder.pack(fill="both", expand=True, padx=12, pady=(0, 12))

        tree = ttk.Treeview(holder, columns=("source", "display", "color"), show="headings", selectmode="browse")
        tree.heading("source", text="Source")
        tree.heading("display", text="Nom affiché")
        tree.heading("color", text="Couleur bulle")
        tree.column("source", width=360, stretch=True)
        tree.column("display", width=280, stretch=True)
        tree.column("color", width=140, stretch=False)

        vsb = ttk.Scrollbar(holder, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        sender_col = self.col_sender_var.get()
        participants: list[str] = []
        if sender_col and sender_col in self.df.columns:
            raw = self.df[sender_col].fillna("").astype(str).unique()
            for v in raw:
                participants.append(_normalize_participant(v))
            participants = sorted(set(participants))

        self.display_name_map.setdefault("", self.empty_sender_name_var.get().strip() or "Unknown")
        for p in participants:
            self.display_name_map.setdefault(p, p)

        for p in participants:
            src = p
            disp = self.display_name_map.get(p, p if p else self.display_name_map[""])
            col = self.participant_color_map.get(p, "")
            tree.insert("", "end", values=(src, disp, col))

        edit = ctk.CTkFrame(win, corner_radius=12)
        edit.pack(fill="x", padx=12, pady=(0, 12))

        src_var = tk.StringVar(value="")
        disp_var = tk.StringVar(value="")
        color_var = tk.StringVar(value="")

        def on_select(_):
            sel = tree.selection()
            win._current_iid = sel[0] if sel else None
            if not sel:
                return
            vals = tree.item(sel[0], "values")
            src, disp = vals[0], vals[1]
            col = vals[2] if len(vals) > 2 else ""
            src_var.set(src)
            disp_var.set(disp)
            color_var.set(col)

        tree.bind("<<TreeviewSelect>>", on_select)

        grid = ctk.CTkFrame(edit, fg_color="transparent")
        grid.pack(fill="x", padx=12, pady=10)
        grid.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(grid, text="Source").grid(row=0, column=0, sticky="w", padx=(0, 10), pady=6)
        ctk.CTkEntry(grid, textvariable=src_var, state="disabled").grid(row=0, column=1, sticky="ew", pady=6)

        ctk.CTkLabel(grid, text="Nom affiché").grid(row=1, column=0, sticky="w", padx=(0, 10), pady=6)
        e2 = ctk.CTkEntry(grid, textvariable=disp_var)
        e2.grid(row=1, column=1, sticky="ew", pady=6)

        ctk.CTkLabel(grid, text="Couleur bulle").grid(row=2, column=0, sticky="w", padx=(0, 10), pady=6)
        color_preview = ctk.CTkLabel(grid, textvariable=color_var, text_color="#A8A8A8")
        color_preview.grid(row=2, column=1, sticky="w", pady=6)

        def apply_row(iid=None, disp=None, col=None, allow_empty_disp=False):
            if iid is None:
                iid = getattr(win, "_current_iid", None)
            if not iid:
                sel = tree.selection()
                iid = sel[0] if sel else None
            if not iid:
                return

            vals = tree.item(iid, "values")
            src = vals[0] if vals else ""
            src = str(src or "")
            d = (disp if disp is not None else disp_var.get()).strip()
            if (not d) and (not allow_empty_disp):
                return
            if d:
                self.display_name_map[src] = d

            c = (col if col is not None else color_var.get()).strip()
            if c:
                self.participant_color_map[src] = c
            elif src in self.participant_color_map:
                del self.participant_color_map[src]

            disp_show = self.display_name_map.get(src, src)
            tree.item(iid, values=(src, disp_show, self.participant_color_map.get(src, "")))

        def pick_color():
            init = color_var.get().strip() or "#3D7DFF"
            c = colorchooser.askcolor(initialcolor=init, title="Couleur de la bulle")
            if c and c[1]:
                color_var.set(c[1])
                apply_row()

        ctk.CTkButton(grid, text="Choisir…", width=110, command=pick_color).grid(row=2, column=1, sticky="e", pady=6)

        # Apply on focus-out (no delayed autosave)
        e2.bind("<FocusOut>", lambda e: apply_row())

        def reset_selected():
            sel = tree.selection()
            if not sel:
                messagebox.showwarning("Attention", "Sélectionnez une ligne à réinitialiser.")
                return
            vals = tree.item(sel[0], "values")
            src = vals[0] if vals else src_var.get()
            src = str(src or "")
            src_var.set(src)

            if not messagebox.askokcancel("Confirmer", "Réinitialiser ce participant (nom + couleur) ?"):
                return

            disp = self.empty_sender_name_var.get().strip() or "Unknown" if src == "" else src
            self.display_name_map[src] = disp
            if src in self.participant_color_map:
                del self.participant_color_map[src]
            disp_var.set(disp)
            color_var.set("")
            tree.item(sel[0], values=(src, disp, ""))

        def ok_close():
            apply_row()
            win.destroy()

        btn_bar = ctk.CTkFrame(win, fg_color="transparent")
        btn_bar.pack(fill="x", padx=12, pady=(0, 12))

        ctk.CTkButton(btn_bar, text="Réinitialiser", fg_color="transparent", border_width=1, command=reset_selected).pack(side="left")
        ctk.CTkButton(btn_bar, text="OK", command=ok_close, width=120).pack(side="right")
        ctk.CTkButton(btn_bar, text="Fermer", fg_color="transparent", border_width=1, command=win.destroy).pack(side="right", padx=(8, 0))

        win.bind("<Return>", lambda e: ok_close())
        tree.bind("<Double-1>", lambda e: ok_close())
        win.bind("<Delete>", lambda e: reset_selected())
        tree.bind("<Delete>", lambda e: reset_selected())

    # ---------------------------
    # Status editor
    # ---------------------------
    def open_status_editor(self):
        win = ctk.CTkToplevel(self)
        win.title("Statuts — mapping valeur → libellé")
        win.geometry("780x520")
        win.minsize(720, 480)
        win.grab_set()

        ctk.CTkLabel(win, text="Mapping de statuts", font=ctk.CTkFont(size=15, weight="bold")).pack(
            anchor="w", padx=12, pady=(12, 8)
        )

        holder = ctk.CTkFrame(win)
        holder.pack(fill="both", expand=True, padx=12, pady=(0, 12))

        tree = ttk.Treeview(holder, columns=("value", "label"), show="headings")
        tree.heading("value", text="Valeur (dans Excel)")
        tree.heading("label", text="Libellé (dans Word)")
        tree.column("value", width=360, stretch=True)
        tree.column("label", width=360, stretch=True)

        vsb = ttk.Scrollbar(holder, orient="vertical", command=tree.yview)
        tree.configure(yscrollcommand=vsb.set)
        tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        for k, v in sorted(self.status_map.items(), key=lambda x: str(x[0])):
            tree.insert("", "end", values=(k, v))

        edit = ctk.CTkFrame(win, corner_radius=12)
        edit.pack(fill="x", padx=12, pady=(0, 12))

        val_var = tk.StringVar(value="")
        lab_var = tk.StringVar(value="")

        def on_select(_):
            sel = tree.selection()
            if not sel:
                return
            k, v = tree.item(sel[0], "values")
            val_var.set(k)
            lab_var.set(v)

        tree.bind("<<TreeviewSelect>>", on_select)

        grid = ctk.CTkFrame(edit, fg_color="transparent")
        grid.pack(fill="x", padx=12, pady=10)
        grid.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(grid, text="Valeur").grid(row=0, column=0, sticky="w", padx=(0, 10), pady=6)
        ctk.CTkEntry(grid, textvariable=val_var).grid(row=0, column=1, sticky="ew", pady=6)

        ctk.CTkLabel(grid, text="Libellé").grid(row=1, column=0, sticky="w", padx=(0, 10), pady=6)
        ctk.CTkEntry(grid, textvariable=lab_var).grid(row=1, column=1, sticky="ew", pady=6)

        def add_or_update():
            k = val_var.get().strip()
            v = lab_var.get().strip()
            if not k:
                messagebox.showwarning("Attention", "La valeur ne peut pas être vide.")
                return
            self.status_map[k] = v or k

            sel = tree.selection()
            if sel:
                tree.item(sel[0], values=(k, self.status_map[k]))
            else:
                tree.insert("", "end", values=(k, self.status_map[k]))

        def delete_selected():
            sel = tree.selection()
            if not sel:
                return
            k, _ = tree.item(sel[0], "values")
            if k in self.status_map:
                del self.status_map[k]
            tree.delete(sel[0])

        btns = ctk.CTkFrame(edit, fg_color="transparent")
        btns.pack(fill="x", padx=12, pady=(0, 12))
        ctk.CTkButton(btns, text="Ajouter / Mettre à jour", command=add_or_update).pack(side="left")
        ctk.CTkButton(btns, text="Supprimer", fg_color="transparent", border_width=1, command=delete_selected).pack(
            side="left", padx=10
        )

        def ok_close():
            if val_var.get().strip():
                add_or_update()
            win.destroy()

        bar = ctk.CTkFrame(win, fg_color="transparent")
        bar.pack(fill="x", padx=12, pady=(0, 12))
        ctk.CTkButton(bar, text="OK", command=ok_close, width=120).pack(side="right")
        ctk.CTkButton(bar, text="Fermer", fg_color="transparent", border_width=1, command=win.destroy).pack(
            side="right", padx=(8, 0)
        )

        win.bind("<Return>", lambda e: ok_close())
        tree.bind("<Delete>", lambda e: delete_selected())
        tree.bind("<Double-1>", lambda e: ok_close())

    # ---------------------------
    # Export
    # ---------------------------
    def export_word(self):
        if self.df is None or self.df.empty:
            messagebox.showwarning("Attention", "Aucune donnée chargée.")
            return

        out_path = filedialog.asksaveasfilename(
            title="Enregistrer le DOCX compatible LibreOffice",
            defaultextension=".docx",
            filetypes=[("DOCX Word / LibreOffice Writer", "*.docx")]
        )
        if not out_path:
            return

        col_sender = self.col_sender_var.get()
        col_msg = self.col_msg_var.get()
        col_dir = self.col_dir_var.get().strip() or None

        ts_single = ts_date = ts_time = None
        if self.ts_mode_var.get() == "single":
            ts_single = self.ts_single_var.get().strip() or None
        else:
            ts_date = self.ts_date_var.get().strip() or None
            ts_time = self.ts_time_var.get().strip() or None

        mode_side = self.side_mode_var.get()

        right_participant = None
        owner_col = None
        if mode_side == "sender":
            owner_col = self.owner_col_var.get().strip() or None
            selected_label = self.right_participant_var.get()
            right_participant = self._participant_display_to_value.get(selected_label, selected_label)

        media_col = self.media_col_var.get().strip() or None
        empty_sender_display = self.empty_sender_name_var.get().strip() or "Unknown"
        bubble_width = float(self.bubble_width_var.get())
        gutter_cm = float(self.gutter_cm_var.get())
        bubble_spacing = float(self.bubble_spacing_var.get())

        # Intro text
        try:
            intro_text = self._intro_textbox.get("1.0", "end").strip()
        except Exception:
            intro_text = self.intro_text_var.get().strip()

        # Dedupe uses normalized timestamp series
        ts_series = self.df.apply(lambda row: build_timestamp(row, ts_single, ts_date, ts_time), axis=1)

        df_to_export = self.df
        removed_dupes = 0
        if self.dedupe_var.get():
            try:
                before_n = int(len(df_to_export))
                df_to_export = dedupe_messages_option_b(df_to_export, col_sender, col_msg, ts_series)
                after_n = int(len(df_to_export))
                removed_dupes = max(0, before_n - after_n)
            except Exception as e:
                messagebox.showwarning("Déduplication", f"Déduplication ignorée (erreur):\n{e}")

        if removed_dupes > 0:
            messagebox.showinfo(
                "Doublons supprimés",
                f"{removed_dupes} ligne(s) supprimée(s) (doublons).\nLignes restantes: {len(df_to_export)}"
            )

        status_col = self.status_col_var.get().strip() or None
        status_map = self.status_map if self.show_status_var.get() else {}
        status_col = status_col if self.show_status_var.get() else None

        try:
            export_chat_to_docx(
                df=df_to_export,
                out_path=out_path,
                col_sender=col_sender,
                col_message=col_msg,
                col_direction=col_dir,
                mode_side=mode_side,
                right_participant_value=right_participant,
                owner_col=owner_col,
                ts_single_col=ts_single,
                ts_date_col=ts_date,
                ts_time_col=ts_time,
                display_name_map=self.display_name_map,
                empty_sender_display=empty_sender_display,
                participant_color_map=self.participant_color_map,
                default_left_color=self.left_color,
                default_right_color=self.right_color,
                bubble_width_cm=bubble_width,
                gutter_cm=gutter_cm,
                bubble_spacing_pt=bubble_spacing,
                extra_cols=(self.extra_cols_selected if self.extra_cols_selected else None),
                media_path_col=media_col,
                files_root=self.files_root_var.get().strip(),
                search_subfolders=bool(self.search_subfolders_var.get()),
                fuzzy_match=bool(self.fuzzy_match_var.get()),
                image_width_mode=self.image_width_mode_var.get(),
                image_max_width_cm=float(self.image_max_width_cm_var.get()),
                compress_images=bool(self.compress_images_var.get()),
                jpeg_quality=int(self.jpeg_quality_var.get()),
                max_image_px=int(self.max_image_px_var.get()),
                keep_media_cache=bool(self.keep_media_cache_var.get()),
                video_thumbs=bool(self.video_thumb_var.get()),
                video_thumb_width_cm=float(self.video_thumb_width_cm_var.get()),
                status_col=status_col,
                status_map=status_map,
                orientation=self.orientation_var.get(),
                page_title=self.page_title_var.get(),
                intro_text=intro_text,
                page_numbers=bool(self.page_numbers_var.get()),
                font_name=self.font_name_var.get().strip() or "Calibri",
                font_color_hex=self.font_color_var.get().strip() or "#000000",
                sender_font_size=int(self.sender_font_size_var.get()),
                msg_font_size=int(self.msg_font_size_var.get()),
                meta_font_size=int(self.meta_font_size_var.get()),
            )
        except Exception as e:
            messagebox.showerror("Erreur export", f"Échec de l'export:\n{e}")
            return

        # Sauvegarde automatique après un export réussi : les réglages courants
        # seront retrouvés au prochain lancement, même sans clic manuel.
        self.save_settings(silent=True)

        if self.open_word_var.get():
            open_in_word(out_path)

        messagebox.showinfo("Terminé", f"Export DOCX compatible LibreOffice créé:\n{out_path}")


if __name__ == "__main__":
    app = App()
    app.mainloop()